"""Fargate worker that creates and emails an administrator's requested report."""
import csv
from datetime import datetime, timezone
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
import json
import logging
import os
from pathlib import Path
import sys
import tempfile
import textwrap
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import boto3

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "../.."))

from generic_dals.report_dal import ReportDAL
from services.generate_report.report_validation import validate_task_event
from utils.db import get_db_connection


logging.basicConfig(level=logging.INFO)
log = logging.getLogger(__name__)

REPORT_BATCH_SIZE = 500
REPORT_TIMESTAMP_TIMEZONE_SETTING = "report_timestamp_timezone"
REPORT_COLUMNS = (
    ("Submitted", "submitted_at"),
    ("Student", "student_name"),
    ("Student number", "student_number"),
    ("Email", "student_email"),
    ("Program", "program_initial"),
    ("Year level", "year_level"),
    ("Assessment result", "scenario_name"),
    ("Counseling status", "counseling_status"),
)


def load_task_event(raw_value=None) -> dict:
    raw_value = os.environ.get("TASK_EVENT") if raw_value is None else raw_value
    if not raw_value:
        raise ValueError("TASK_EVENT is required")
    try:
        return validate_task_event(json.loads(raw_value))
    except (TypeError, ValueError, json.JSONDecodeError) as exc:
        raise ValueError("TASK_EVENT is invalid") from exc


def iter_report_rows(dal, filters: dict, batch_size=REPORT_BATCH_SIZE):
    """Yield report rows without holding the complete assessment history in memory."""
    cursor = None
    while True:
        rows = dal.get_report_batch(filters, batch_size, cursor)
        if not rows:
            return
        for row in rows:
            yield dict(row)
        last_row = rows[-1]
        cursor = (last_row["submitted_at"], last_row["assessment_id"])
        if len(rows) < batch_size:
            return


def _plain_text(value, display_timezone=timezone.utc) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.astimezone(display_timezone).strftime("%Y-%m-%d %H:%M %Z")
    return " ".join(str(value).replace("\x00", "").split())


def _spreadsheet_text(value, display_timezone=timezone.utc) -> str:
    value = _plain_text(value, display_timezone)
    # Student-entered strings must not become spreadsheet formulas on export.
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


def _row_values(row: dict, display_timezone=timezone.utc) -> list[str]:
    return [_spreadsheet_text(row.get(key), display_timezone) for _, key in REPORT_COLUMNS]


def get_report_timezone(dal) -> tuple[ZoneInfo, str]:
    """Read and validate the server-controlled report timestamp timezone once."""
    configured_timezone = dal.get_setting(REPORT_TIMESTAMP_TIMEZONE_SETTING)
    if configured_timezone is None:
        raise RuntimeError("The report timestamp timezone setting is missing.")
    if not isinstance(configured_timezone, str) or not configured_timezone.strip():
        raise RuntimeError("The report timestamp timezone setting is invalid.")

    timezone_name = configured_timezone.strip()
    try:
        return ZoneInfo(timezone_name), timezone_name
    except ZoneInfoNotFoundError as exc:
        raise RuntimeError("The report timestamp timezone setting is invalid.") from exc


def _metadata_rows(task_event: dict, generated_at: datetime | None = None, timezone_name: str | None = None) -> list[list[str]]:
    filters = task_event["filters"]
    filter_summary = []
    for name in ("programs", "years", "counseling_status_ids", "scenario_ids", "start_date", "end_date"):
        if filters.get(name):
            value = filters[name]
            filter_summary.append(f"{name}: {', '.join(map(str, value)) if isinstance(value, list) else value}")
    return [
        ["Report type", task_event["report_type"].title()],
        ["Generated at", (generated_at or datetime.now(timezone.utc)).strftime("%Y-%m-%d %H:%M:%S %Z")
         + (f" ({timezone_name})" if timezone_name else "")],
        ["Filters", "; ".join(filter_summary) or "All matching assessments"],
        ["Recommendations", task_event["filters"].get("recommendations", "")],
    ]


def report_filename(generated_at: datetime, report_format: str) -> str:
    """Create a timestamped, filesystem-safe attachment filename."""
    timestamp = generated_at.strftime("%Y%m%dT%H%M%S%z")
    return f"imhealth-report-{timestamp}.{report_format}"


def _write_csv(path: Path, rows, task_event: dict, generated_at=None, timezone_name=None, display_timezone=timezone.utc) -> int:
    count = 0
    with path.open("w", encoding="utf-8", newline="") as output:
        writer = csv.writer(output)
        for row in _metadata_rows(task_event, generated_at, timezone_name):
            writer.writerow([_spreadsheet_text(value) for value in row])
        writer.writerow([])
        writer.writerow([column for column, _ in REPORT_COLUMNS])
        for row in rows:
            writer.writerow(_row_values(row, display_timezone))
            count += 1
    return count


def _write_xlsx(path: Path, rows, task_event: dict, generated_at=None, timezone_name=None, display_timezone=timezone.utc) -> int:
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font
    except ImportError as exc:
        raise RuntimeError("The XLSX report dependency is unavailable.") from exc

    workbook = Workbook(write_only=True)
    metadata = workbook.create_sheet("Report details")
    for row in _metadata_rows(task_event, generated_at, timezone_name):
        metadata.append([_spreadsheet_text(value) for value in row])
    sheet = workbook.create_sheet("Assessments")
    header = []
    for column, _ in REPORT_COLUMNS:
        cell = WriteOnlyCell(sheet, value=column)
        cell.font = Font(bold=True)
        header.append(cell)
    sheet.append(header)
    count = 0
    for row in rows:
        sheet.append(_row_values(row, display_timezone))
        count += 1
    workbook.save(path)
    return count


def _write_pdf(path: Path, rows, task_event: dict, generated_at=None, timezone_name=None, display_timezone=timezone.utc) -> int:
    try:
        from reportlab.lib.pagesizes import letter, landscape
        from reportlab.pdfgen import canvas
    except ImportError as exc:
        raise RuntimeError("The PDF report dependency is unavailable.") from exc

    page_size = landscape(letter)
    document = canvas.Canvas(str(path), pagesize=page_size)
    page_width, page_height = page_size
    left = 36
    y = page_height - 36

    def heading(include_metadata=True):
        nonlocal y
        document.setFont("Helvetica-Bold", 15)
        document.drawString(left, y, f"PUP iMHealth {task_event['report_type'].title()} Report")
        y -= 20
        if include_metadata:
            document.setFont("Helvetica", 8)
            for label, value in _metadata_rows(task_event, generated_at, timezone_name):
                for line in textwrap.wrap(f"{label}: {_plain_text(value)}", width=150) or [""]:
                    if y < 48:
                        document.showPage()
                        y = page_height - 36
                        document.setFont("Helvetica-Bold", 12)
                        document.drawString(left, y, "PUP iMHealth Report details (continued)")
                        y -= 18
                        document.setFont("Helvetica", 8)
                    document.drawString(left, y, line)
                    y -= 11
            if y < 60:
                document.showPage()
                y = page_height - 36
        y -= 4
        document.setFont("Helvetica-Bold", 8)
        document.drawString(left, y, "Submitted | Student | Program | Year | Result | Counseling status")
        y -= 12
        document.setFont("Helvetica", 8)

    heading()
    count = 0
    for row in rows:
        if y < 40:
            document.showPage()
            y = page_height - 36
            heading(include_metadata=False)
        line = " | ".join((
            _plain_text(row.get("submitted_at"), display_timezone),
            _plain_text(row.get("student_name"), display_timezone),
            _plain_text(row.get("program_initial"), display_timezone),
            _plain_text(row.get("year_level"), display_timezone),
            _plain_text(row.get("scenario_name"), display_timezone),
            _plain_text(row.get("counseling_status"), display_timezone),
        ))
        document.drawString(left, y, textwrap.shorten(line, width=175, placeholder="…"))
        y -= 12
        count += 1
    if count == 0:
        document.drawString(left, y, "No assessments matched the selected filters.")
    document.save()
    return count


def write_report(dal, task_event: dict, directory=None) -> tuple[Path, int]:
    """Stream database batches into the requested report format on local task storage."""
    directory = directory or tempfile.gettempdir()
    report_timezone, timezone_name = get_report_timezone(dal)
    generated_at = datetime.now(report_timezone)
    path = Path(directory) / report_filename(generated_at, task_event["format"])
    path.touch(exist_ok=False)
    rows = iter_report_rows(dal, task_event["filters"], batch_size=REPORT_BATCH_SIZE)
    writers = {"csv": _write_csv, "xlsx": _write_xlsx, "pdf": _write_pdf}
    try:
        count = writers[task_event["format"]](
            path, rows, task_event, generated_at, timezone_name, report_timezone
        )
        return path, count
    except Exception:
        path.unlink(missing_ok=True)
        raise


def send_report_email(ses, from_email: str, recipient_email: str, report_path: Path, row_count: int):
    message = MIMEMultipart()
    message["Subject"] = "Your requested PUP iMHealth report"
    message["From"] = from_email
    message["To"] = recipient_email
    message.attach(MIMEText(
        f"Your requested report is attached. It contains {row_count} matching assessment record(s).",
        "plain",
        "utf-8",
    ))
    attachment = MIMEApplication(report_path.read_bytes())
    attachment.add_header("Content-Disposition", "attachment", filename=report_path.name)
    message.attach(attachment)
    ses.send_raw_email(
        Source=from_email,
        Destinations=[recipient_email],
        RawMessage={"Data": message.as_string()},
    )


def main():
    task_event = load_task_event()
    report_path = None
    conn = None
    try:
        conn = get_db_connection()
        report_path, row_count = write_report(ReportDAL(conn), task_event)
        from_email = os.environ["SES_FROM_EMAIL"]
        ses = boto3.client("ses", region_name=os.environ.get("AWS_REGION", "ap-southeast-1"))
        send_report_email(ses, from_email, task_event["recipient_email"], report_path, row_count)
        log.info("Sent %s %s report with %s rows", task_event["report_type"], task_event["format"], row_count)
    finally:
        if report_path:
            report_path.unlink(missing_ok=True)
        if conn:
            conn.close()


if __name__ == "__main__":
    main()
